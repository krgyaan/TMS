"""
Test multi-page document to verify header and footer repetition.
"""
from pathlib import Path
import fitz
import win32com.client
from docx import Document
from docx.shared import Inches, Pt

doc = Document()
sec = doc.sections[0]
sec.page_width = Pt(595.3)
sec.page_height = Pt(841.9)
sec.left_margin = Inches(0.5)
sec.right_margin = Inches(0.5)
sec.top_margin = Pt(128)
sec.bottom_margin = Pt(65)
sec.header_distance = Pt(0)
sec.footer_distance = Pt(32)

assets_dir = Path("assets")
# Header
h = sec.header.paragraphs[0]
h.paragraph_format.space_before = Pt(0)
h.paragraph_format.space_after = Pt(0)
h.paragraph_format.left_indent = -sec.left_margin
h.paragraph_format.right_indent = -sec.right_margin
h.add_run().add_picture(str(assets_dir / "header_band.png"), width=sec.page_width)

# Footer
f = sec.footer.paragraphs[0]
f.paragraph_format.space_before = Pt(0)
f.paragraph_format.space_after = Pt(0)
f.paragraph_format.left_indent = -sec.left_margin
f.paragraph_format.right_indent = -sec.right_margin
f.add_run().add_picture(str(assets_dir / "footer_band.png"), width=sec.page_width)

# Add 2 pages of content
for page_num in range(1, 3):
    doc.add_heading(f"Page {page_num} Title", level=1)
    for i in range(25):
        doc.add_paragraph(f"Paragraph {i+1} on page {page_num} with sample text.")
    if page_num < 2:
        doc.add_page_break()

test_multi = assets_dir / "test_multi.docx"
doc.save(str(test_multi))

word = win32com.client.Dispatch("Word.Application")
word.Visible = False
try:
    d = word.Documents.Open(str(test_multi.resolve()))
    pdf_out = str((assets_dir / "test_multi.pdf").resolve())
    d.SaveAs(pdf_out, FileFormat=17)
    d.Close(False)
finally:
    word.Quit()

p_doc = fitz.open(pdf_out)
print("Multi-page test: Page count =", len(p_doc))
for i, page in enumerate(p_doc):
    imgs = page.get_image_info()
    print(f"Page {i+1} image count: {len(imgs)}")
    for img in imgs:
        print(f"  bbox: {img['bbox']}")
