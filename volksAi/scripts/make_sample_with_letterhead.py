"""
One-off: build the Format F-2A fixture with the Volks letterhead applied, save it as
assets/sample_with_letterhead.docx, then verify it (python-docx read-back + Word->PDF
render + overlap check). Does NOT change build_annexure_docx()'s default behaviour.

Run from volksAi/:  python scripts/make_sample_with_letterhead.py
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import fitz
import win32com.client
from docx import Document

from app.services.annexure_docx_builder import (
    TOP_MARGIN_PT, BOTTOM_MARGIN_PT, HEADER_BAND_HEIGHT_PT, apply_letterhead, build_annexure_docx,
)
from app.services.test_annexures import F2A_ANNEXURE

OUT = ROOT / "assets" / "sample_with_letterhead.docx"
PDF = ROOT / "assets" / "sample_with_letterhead.pdf"

# build_annexure_docx() saves the file itself, so apply the letterhead on a re-open.
build_annexure_docx(F2A_ANNEXURE, OUT)
doc = Document(str(OUT))
apply_letterhead(doc)
doc.save(str(OUT))
print("saved", OUT)

# --- read-back -------------------------------------------------------------
back = Document(str(OUT))
sec = back.sections[0]
for name, part in (("header", sec.header), ("footer", sec.footer)):
    pics = part._element.xpath(".//pic:pic")
    print(f"{name}: {len(pics)} picture(s); related images:",
          [r.target_ref for r in part.part.rels.values() if "image" in r.reltype])
print(f"page {sec.page_width.pt:.1f} x {sec.page_height.pt:.1f} pt; top margin {sec.top_margin.pt} pt; "
      f"bottom margin {sec.bottom_margin.pt} pt; body paragraphs {len(back.paragraphs)}")

# --- render through Word and measure ----------------------------------------
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
try:
    d = word.Documents.Open(str(OUT.resolve()))
    d.SaveAs(str(PDF.resolve()), FileFormat=17)
    d.Close(False)
finally:
    word.Quit()

pdf = fitz.open(str(PDF))
print("rendered pages:", len(pdf))
for i, page in enumerate(pdf):
    H = page.rect.height
    bands = [fitz.Rect(im["bbox"]) for im in page.get_image_info()]
    hdr = [b for b in bands if b.y0 < H / 2]
    ftr = [b for b in bands if b.y0 >= H / 2]
    words = page.get_text("words")
    body_top = min(w[1] for w in words if w[1] > HEADER_BAND_HEIGHT_PT)
    body_bot = max(w[3] for w in words if w[3] < H - 0)  # filtered below
    # body text = words that are not inside an image band
    body = [w for w in words if not any(fitz.Rect(w[:4]).intersects(b) for b in bands)]
    b_top, b_bot = min(w[1] for w in body), max(w[3] for w in body)
    hdr_bottom = max(b.y1 for b in hdr); ftr_top = min(b.y0 for b in ftr)
    print(f"p{i+1}: header img y {hdr[0].y0:.1f}-{hdr_bottom:.1f} | footer img y {ftr_top:.1f}-{max(b.y1 for b in ftr):.1f} "
          f"| body text y {b_top:.1f}-{b_bot:.1f} | gap above {b_top-hdr_bottom:.1f} pt, gap below {ftr_top-b_bot:.1f} pt "
          f"| overlap: {b_top < hdr_bottom or b_bot > ftr_top}")
    page.get_pixmap(dpi=110).save(str(ROOT / "assets" / f"sample_with_letterhead_p{i+1}.png"))
