"""
Tests for the broadened Role 3 "tender knowledge" pass: one Sonnet read of the page-tagged
main+ATC text now returns requirements[] AND annexures[] (additive -- requirements are
unchanged), the endpoint reports schemaVersion, and a separate deterministic
POST /generate-annexure-docx renders one stored annexure without any LLM call.

Anthropic is mocked throughout; no live API calls. The annexure fixture reproduces real
GAIL "Format F-2A Declaration for Bid Security" phrasing (gold-standard GAIL ATC).
"""
import io
import json
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.bidding_requirements_resolver import (
    SYSTEM_PROMPT,
    TENDER_KNOWLEDGE_SCHEMA_VERSION,
    _build_requirements_tool_schema,
    analyze_bidding_requirements,
)

client = TestClient(app)

TENDER_TEXT = (
    "[Main Page 3]: Bidder must submit OEM Authorization Certificate. Bidders are required to submit "
    "Declaration for Bid Security in bid as per proforma at Form F-2A.\n\n"
    "[Main Page 68]: FORMAT F-2A DECLARATION FOR BID SECURITY To, M/s GAIL (INDIA) LIMITED "
    "SUB: TENDER NO: we M/s____________ (Name of Bidder) have submitted our offer/ bid no. ....... "
    "Place: [Signature of Authorized Signatory of Bidder] Date: Name: Designation:"
)

REQUIREMENT = {
    "documentName": "OEM Authorization Certificate",
    "category": "oem",
    "required": True,
    "source": {"document": "main", "page": 3, "snippet": "Bidder must submit OEM Authorization Certificate"},
    "matchedLibraryId": "lib-7",  # model wrongly matched an OEM item -> guard must null it
    "confidence": "high",
    "reasoning": "Explicit requirement.",
}
F2A_ANNEXURE = {
    "annexureName": "Format F-2A: Declaration for Bid Security",
    "source": {"document": "main", "page": 68, "snippet": "FORMAT F-2A DECLARATION FOR BID SECURITY"},
    "blocks": [
        {"type": "heading", "text": "FORMAT F-2A DECLARATION FOR BID SECURITY"},
        {"type": "paragraph", "text": "we M/s______ (Name of Bidder) have submitted our offer/ bid no. ______"},
        {"type": "blank_field", "label": "Place"},
        {"type": "signature_line", "label": "Signature of Authorized Signatory of Bidder"},
        {"type": "blank_field", "label": "Designation"},
    ],
}


def _mock_response(tool_input, stop_reason="tool_use"):
    tool_block = SimpleNamespace(type="tool_use", name="report_bidding_requirements", input=tool_input)
    usage = SimpleNamespace(input_tokens=900, output_tokens=700, cache_creation_input_tokens=0, cache_read_input_tokens=0)
    return SimpleNamespace(content=[tool_block], usage=usage, stop_reason=stop_reason)


def _analyze(tool_input, **kw):
    with patch("anthropic.Anthropic") as MockAnthropic:
        MockAnthropic.return_value.messages.create.return_value = _mock_response(tool_input, **kw)
        result = analyze_bidding_requirements(TENDER_TEXT, library_documents=[], api_key="sk-ant-test-key")
        call_kwargs = MockAnthropic.return_value.messages.create.call_args.kwargs
    return result, call_kwargs


# ── Schema / prompt are additive ─────────────────────────────────────────────

def test_tool_schema_adds_annexures_without_changing_requirements():
    schema = _build_requirements_tool_schema()
    props = schema["input_schema"]["properties"]
    assert schema["name"] == "report_bidding_requirements"            # tool name unchanged
    assert set(props) == {"requirements", "annexures"}
    assert schema["input_schema"]["required"] == ["requirements", "annexures"]  # both tasks must be answered
    assert props["requirements"]["items"]["required"] == [
        "documentName", "category", "required", "source", "matchedLibraryId", "confidence", "reasoning",
    ]
    block_types = props["annexures"]["items"]["properties"]["blocks"]["items"]["properties"]["type"]["enum"]
    assert block_types == ["heading", "paragraph", "blank_field", "table", "signature_line"]


def test_system_prompt_covers_both_tasks():
    assert "Identify EVERY document or certificate a bidder must submit" in SYSTEM_PROMPT  # task 1 intact
    assert "SECOND TASK" in SYSTEM_PROMPT and "annexure" in SYSTEM_PROMPT.lower()
    assert "Do not invent fields" in SYSTEM_PROMPT
    assert "describe what it is for" in SYSTEM_PROMPT


# ── (a) one call returns both ────────────────────────────────────────────────

def test_single_call_parses_requirements_and_annexures():
    result, call_kwargs = _analyze({"requirements": [REQUIREMENT], "annexures": [F2A_ANNEXURE]})

    assert call_kwargs["tool_choice"] == {"type": "tool", "name": "report_bidding_requirements"}
    assert len(result["requirements"]) == 1
    assert result["requirements"][0]["matchedLibraryId"] is None  # existing OEM guard still applies
    assert [a["annexureName"] for a in result["annexures"]] == ["Format F-2A: Declaration for Bid Security"]
    assert result["annexures"][0]["blocks"] == F2A_ANNEXURE["blocks"]
    assert result["rejectedAnnexures"] == []
    assert result["truncated"] is False
    assert result["usage"]["input_tokens"] == 900


def test_claude_client_uses_role3_timeout_sized_for_both_tasks():
    from app.services.bidding_requirements_resolver import ROLE_3_TIMEOUT_S
    assert ROLE_3_TIMEOUT_S == 180.0  # was 60s when the call returned requirements only
    with patch("anthropic.Anthropic") as MockAnthropic:
        MockAnthropic.return_value.messages.create.return_value = _mock_response({"requirements": [], "annexures": []})
        analyze_bidding_requirements(TENDER_TEXT, library_documents=[], api_key="sk-ant-test-key")
    assert MockAnthropic.call_args.kwargs["timeout"] == 180.0


# ── (b) old-shape response ───────────────────────────────────────────────────

def test_old_shape_response_without_annexures_defaults_to_empty():
    result, _ = _analyze({"requirements": [REQUIREMENT]})
    assert len(result["requirements"]) == 1
    assert result["annexures"] == []
    assert result["rejectedAnnexures"] == []


def test_null_annexures_does_not_crash():
    result, _ = _analyze({"requirements": [], "annexures": None})
    assert result["annexures"] == []


def test_truncation_is_flagged():
    result, _ = _analyze({"requirements": [REQUIREMENT]}, stop_reason="max_tokens")
    assert result["truncated"] is True


# ── (d) citation guard ───────────────────────────────────────────────────────

@pytest.mark.parametrize("bad_source", [
    None,                                                        # no source at all
    {"document": "main", "snippet": "FORMAT F-2A"},              # no page
    {"document": "main", "page": 99, "snippet": "FORMAT F-2A"},  # page never in the text
])
def test_citationless_annexure_is_rejected_not_accepted(bad_source):
    bad = {k: v for k, v in F2A_ANNEXURE.items() if k != "source"}
    if bad_source is not None:
        bad["source"] = bad_source
    result, _ = _analyze({"requirements": [REQUIREMENT], "annexures": [bad]})
    assert result["annexures"] == []
    assert len(result["rejectedAnnexures"]) == 1
    assert result["rejectedAnnexures"][0]["annexureName"] == F2A_ANNEXURE["annexureName"]
    assert len(result["requirements"]) == 1  # requirements unaffected by a rejected annexure


# ── Endpoint: schemaVersion + annexures in the response ─────────────────────

def test_endpoint_reports_schema_version_and_annexures():
    mock_result = {
        "requirements": [REQUIREMENT],
        "annexures": [F2A_ANNEXURE],
        "rejectedAnnexures": [{"annexureName": "Bogus", "reason": "missing source citation", "raw": {"x": 1}}],
        "truncated": False,
        "usage": {"input_tokens": 1},
    }
    with patch("app.routers.bidding_requirements.extract_pdf_text_hybrid", return_value=[{"page": 68, "text": "FORMAT F-2A"}]), \
         patch("app.routers.bidding_requirements.analyze_bidding_requirements", return_value=mock_result):
        response = client.post(
            "/analyze-bidding-requirements",
            files={"pdf_file": ("tender.pdf", b"%PDF-1.4", "application/pdf")},
            data={"library_documents": "[]"},
        )
    assert response.status_code == 200
    data = response.json()
    assert data["schemaVersion"] == TENDER_KNOWLEDGE_SCHEMA_VERSION == 1
    assert data["requirements"][0]["documentName"] == "OEM Authorization Certificate"
    assert data["annexures"][0]["annexureName"] == F2A_ANNEXURE["annexureName"]
    assert data["rejectedAnnexures"] == [{"annexureName": "Bogus", "reason": "missing source citation"}]
    assert "docxPath" not in data["annexures"][0]  # identification only -- no file generated here


def test_endpoint_with_old_shape_resolver_result_defaults_annexures():
    with patch("app.routers.bidding_requirements.extract_pdf_text_hybrid", return_value=[{"page": 1, "text": "x"}]), \
         patch("app.routers.bidding_requirements.analyze_bidding_requirements",
               return_value={"requirements": [], "usage": None}):
        data = client.post(
            "/analyze-bidding-requirements", files={"pdf_file": ("t.pdf", b"%PDF", "application/pdf")},
        ).json()
    assert data["schemaVersion"] == 1 and data["annexures"] == [] and data["rejectedAnnexures"] == []


# ── (c) /generate-annexure-docx: deterministic, all 5 block types, no LLM ───

ALL_TYPES = {
    "annexureName": "Annexure-III: Manufacturer's Authorization Form",
    "blocks": [
        {"type": "heading", "text": "ANNEXURE-III MANUFACTURER'S AUTHORIZATION FORM"},
        {"type": "paragraph", "text": "We, the undersigned manufacturer, hereby authorize the bidder."},
        {"type": "blank_field", "label": "Name of Authorized Signatory"},
        {"type": "table", "headers": ["Sr. No.", "Item", "Make/Model"], "rows": [["1", "Ni-Cd Battery Bank", ""]]},
        {"type": "signature_line", "label": "Signature & Seal of Bidder"},
    ],
}


def test_generate_docx_from_stored_blocks_without_any_llm_call():
    with patch("anthropic.Anthropic") as MockAnthropic, \
         patch("app.routers.bidding_requirements.analyze_bidding_requirements") as mock_analyze:
        response = client.post("/generate-annexure-docx", json=ALL_TYPES)
    MockAnthropic.assert_not_called()
    mock_analyze.assert_not_called()

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert 'filename="annexure-iii-manufacturer-s-authorization-form.docx"' in response.headers["content-disposition"]

    doc = Document(io.BytesIO(response.content))  # valid, openable .docx
    texts = [p.text for p in doc.paragraphs]
    assert doc.paragraphs[0].text == "ANNEXURE-III MANUFACTURER'S AUTHORIZATION FORM"
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert "We, the undersigned manufacturer, hereby authorize the bidder." in texts
    assert any(t.startswith("Name of Authorized Signatory: ___") for t in texts)
    assert any("Signature & Seal of Bidder" in t for t in texts)
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [
        ["Sr. No.", "Item", "Make/Model"], ["1", "Ni-Cd Battery Bank", ""],
    ]


def test_generate_docx_is_deterministic():
    a = client.post("/generate-annexure-docx", json=ALL_TYPES).content
    b = client.post("/generate-annexure-docx", json=ALL_TYPES).content
    assert [p.text for p in Document(io.BytesIO(a)).paragraphs] == [p.text for p in Document(io.BytesIO(b)).paragraphs]


@pytest.mark.parametrize("payload", [
    {"annexureName": "X", "blocks": []},
    {"annexureName": "X", "blocks": [{"type": "image"}, {"type": "heading"}]},
])
def test_generate_docx_rejects_payload_with_no_valid_blocks(payload):
    assert client.post("/generate-annexure-docx", json=payload).status_code == 400
